"""Bounded audit execution and structural validation; no model calls on import."""
import contextlib
import fcntl
import json
import os
from pathlib import Path
import re
import shutil
import signal
import subprocess
import threading
import time
from functools import lru_cache
from openpyxl import load_workbook

REQUIRED = [f'{i:02d}' for i in range(1,31)] + [f'M{i:02d}' for i in range(11)] + [f'X{i:02d}' for i in range(1,4)] + [f'E{i:02d}' for i in range(1,11)]
POLICY = '''ZASADY WYKONANIA PLATFORMY (pierwszeństwo nad organizacją pracy w instrukcji):
Wykonujesz audyt bez pośrednika zlecającego cały audyt kolejnemu agentowi.
Zachowaj pełny zakres i źródła. Oszczędzaj tokeny: nie czytaj ponownie poprawnych zapisanych wyników OCR jako obrazów bez konkretnej przyczyny.
OCR: maksymalnie 2 subagentów jednocześnie, po 8–12 stron brakujących lub zmienionych; wynik do pliku, odpowiedź tylko ścieżka i braki. Nie dziedzicz pełnej historii dokumentacji. Nie deleguj całego audytu.
Eksperci E01–E10: analizuj kolejno branżami na zapisanych danych; osobny agent tylko dla konkretnej nierozstrzygniętej rozbieżności. Każda opinia ma źródła. Nie przekazuj wszystkim pełnych JSON-ów wszystkich branż.
Po każdym etapie zapisuj wynik i checkpoint.json (ukończone, brakujące, pliki wynikowe). Indeksuj aktualne dokumenty Drive, porównaj zmiany; nie zakładaj kompletności starego read_log bez sprawdzenia dowodów.
Nie uruchamiaj dodatkowych CLI modeli przez Bash. Nie zmieniaj konfiguracji limitów ani kodu platformy. Nie wysyłaj wiadomości.
Wynik końcowy: arkusze 01–30, M00–M10, X01–X03, E01–E10, bez _TMP; każda zakładka ma treść lub jawne uzasadnienie braku zastosowania. Przelicz formuły. Nie deklaruj ukończenia po zapisaniu części pliku.
'''

def write_json(path, value):
    path=Path(path); path.parent.mkdir(parents=True,exist_ok=True)
    temp=path.with_name(path.name+'.tmp'); temp.write_text(json.dumps(value,ensure_ascii=False,indent=2)); os.replace(temp,path)

@lru_cache(maxsize=128)
def _validate(path, size, mtime_ns):
    errors=[]; missing=[]; formulas=uncached=0
    try:
        with contextlib.closing(load_workbook(path,read_only=True,data_only=False)) as wb, contextlib.closing(load_workbook(path,read_only=True,data_only=True)) as values:
            names={s.title.split('_')[0]:s for s in wb}
            missing=[n for n in REQUIRED if n not in names]
            for n in REQUIRED:
                if n in names:
                    rows=sum(any(c.value is not None for c in row) for row in names[n].iter_rows())
                    if rows<2: errors.append('Pusta zakładka '+n)
            if '_TMP' in wb.sheetnames: errors.append('Pozostawiona zakładka robocza _TMP')
            for ws in wb:
                for row,cached in zip(ws.iter_rows(),values[ws.title].iter_rows()):
                    for c,v in zip(row,cached):
                        if c.data_type=='f':
                            formulas+=1
                            if v.value is None: uncached+=1
                        if v.data_type=='e': errors.append('Błąd formuły '+ws.title+'!'+v.coordinate)
            if uncached: errors.append(f'Nieprzeliczone formuły: {uncached}')
            if missing: errors.append('Brak zakładek: '+', '.join(missing))
            return dict(complete=not errors,errors=errors,missing=missing,sheets=len(wb.sheetnames),formulas=formulas,uncached=uncached)
    except Exception as exc:
        return dict(complete=False,errors=['Nie można odczytać XLSX: '+type(exc).__name__],missing=REQUIRED)

def validate(path):
    path=Path(path)
    try: st=path.stat()
    except OSError: return dict(complete=False,errors=['Brak pliku XLSX'],missing=REQUIRED)
    return _validate(str(path),st.st_size,st.st_mtime_ns)

def workdirs(base,fid):
    matches=[]
    for p in Path(base).glob('*/input_index.json'):
        try:
            data=json.loads(p.read_text()); folder=str(data.get('folder_drive') or data.get('folder_id') or '')
            if folder==fid or re.search(r'(?:folders/|id=)'+re.escape(fid)+r'(?:[/?&#]|$)',folder): matches.append(str(p.parent))
        except (OSError,ValueError,AttributeError): pass
    return sorted(matches)

def status(path):
    marker=Path(str(path)+'.status.json')
    if not marker.exists(): return {'state':'unverified','complete':None}
    result=validate(path)
    try: recorded=json.loads(marker.read_text())
    except (OSError,ValueError): recorded={}
    result=dict(result)
    result['complete']=bool(result['complete'] and recorded.get('state')=='complete')
    result['state']='complete' if result['complete'] else 'partial'
    return result

def run(eng,jid,url,hint=None,source=None,alias=None):
    fid=eng.extract_folder_id(url)
    if not re.fullmatch(r'[A-Za-z0-9_-]{20,}',fid):
        eng.jset(jid,done=True,ok=False,stage='Nieprawidłowy folder'); return
    root=Path(eng.BASE)/'work'/'audit_runs'/fid; root.mkdir(parents=True,exist_ok=True)
    lock=(root/'run.lock').open('a')
    try:
        try: fcntl.flock(lock,fcntl.LOCK_EX|fcntl.LOCK_NB)
        except BlockingIOError:
            eng.jset(jid,done=True,ok=False,stage='Analiza już trwa',log_add='ERR\tTen folder ma już aktywny audyt.'); return
        _run_locked(eng,jid,url,hint,fid,root,source=source,alias=alias)
    except Exception as exc:
        eng.jset(jid,done=True,ok=False,stage='Przerwane',log_add='ERR\tBłąd wykonania: '+type(exc).__name__)
        write_json(root/'last_error.json',dict(job_id=jid,error=type(exc).__name__,time=time.time()))
    finally: lock.close()

def _run_locked(eng,jid,url,hint,fid,root,source=None,alias=None):
    target=root/'AUDYT_WYNIK.xlsx'
    statefile=root/'state.json'
    prior={}
    if statefile.exists():
        try: prior=json.loads(statefile.read_text())
        except ValueError: pass
    if source and (not target.exists() or prior.get('state')=='complete'):
        shutil.copy2(source,target)
    dirs=workdirs(eng.BASE,fid)
    state=dict(job_id=jid,folder_id=fid,workdirs=dirs,target=str(target),state='running',started=time.time(),attempts=[],previous_job=prior.get('job_id'))
    write_json(statefile,state)
    instruction=root/'instructions.md'
    instruction.write_text(Path(eng.AGENT_MD_SRC).read_text()+'\n\n'+POLICY)
    prompt=(f'Wykonaj pełny audyt folderu {url}. Instrukcja merytoryczna: {instruction}. '
            f'Najpierw przeczytaj instrukcję, potem zastosuj poniższe zasady wykonania.\n{POLICY}\n'
            f'Zapisuj dane robocze w {root}; trwałe wyniki poprzednich etapów: {dirs}. '
            f'Jeżeli istnieje {target}, kontynuuj ten częściowy plik. Wykorzystaj istniejące generatory i dane; sprawdź ich kompletność. '
            f'JEDYNY plik wynikowy tej próby: {target}. Nie zapisuj ani nie modyfikuj plików w {eng.OUTPUTS}; publikację wykona serwer po walidacji. '
            f'Budżet Claude {os.environ.get("AUDYT_MAX_BUDGET_USD","10")} USD wg CLI i ograniczony czas; zapisuj postęp na bieżąco. '
            f'Wskazówka operatora: {hint or "brak"}')
    selected,cfg=eng.engine_config(alias) if alias else eng.engine_config()
    # alias z routera (model=gemini|gpt) = pierwszy silnik wprost; fallback jak dotąd
    aliases=[selected]+[a for a in eng.FALLBACK_CHAIN if a!=selected][:1]
    deadline=time.monotonic()+float(os.environ.get('AUDYT_TIMEOUT_SECONDS','1800'))
    final_ok=False
    for index,alias in enumerate(aliases):
        if index and not eng.provider_available(alias): break
        remaining=deadline-time.monotonic()
        if remaining<=0: break
        eng.jset(jid,stage='Analiza: '+alias,pct=5,log_add='INFO\tWznawiam zapisane dane; silnik '+alias)
        result=_execute(eng,jid,alias,prompt,remaining)
        check=validate(target)
        state['attempts'].append(dict(provider=alias,**result,validation=check)); write_json(statefile,state)
        eng.jset(jid,usage=result.get('model_usage'),cost_usd=result.get('cost_usd'),validation=check)
        if result.get('cost_usd') is not None:
            eng.jset(jid,log_add=f"INFO\tLicznik CLI ({alias}): {result['cost_usd']:.2f} USD; wycena użycia, nie faktura.")
        if result['ok'] and check['complete']:
            final_ok=True; break
        if result.get('budget_hit') or result.get('timed_out'): break
        # Only quota permits an automatic provider switch, including when XLSX is partial.
        if not result.get('limit_hit'): break
        prompt+='\nPoprzedni silnik wyczerpał limit. Kontynuuj zapisany plik i checkpointy. Braki: '+json.dumps(check,ensure_ascii=False)
    if final_ok:
        dest=Path(eng.OUTPUTS)/f'AUDYT_{fid}_{jid}.xlsx'
        temporary=dest.with_suffix('.xlsx.tmp'); shutil.copy2(target,temporary); os.replace(temporary,dest)
        write_json(str(dest)+'.status.json',dict(state='complete',job_id=jid,folder_id=fid))
        state.update(state='complete',result_id=eng.audit_id(str(dest)))
        eng.jset(jid,done=True,ok=True,pct=100,stage='Gotowe',result_id=state['result_id'],log_add='OK\tKomplet arkuszy i wyniki formuł sprawdzone.')
    else:
        state['state']='partial'
        last=state['attempts'][-1] if state['attempts'] else {}
        reason='Limit budżetu' if last.get('budget_hit') else ('Limit czasu' if last.get('timed_out') else 'Audyt nieukończony')
        eng.jset(jid,done=True,ok=False,pct=95,stage=reason,result_id=None,partial=True,log_add='ERR\t'+reason+'. Dane zachowane. Kolejne uruchomienie wykorzysta zapisane etapy. '+ '; '.join(validate(target)['errors'])[:450])
    state['finished']=time.time();write_json(statefile,state)
    write_json(root/(jid+'.json'),state)

def _execute(eng,jid,alias,prompt,timeout):
    cmd=eng.engine_cmd(prompt,jid,alias)
    if eng.engine_config(alias)[1]['provider']=='claude': cmd+=['--agent','audyt-przetargowy','--append-system-prompt',POLICY]
    proc=subprocess.Popen(cmd,cwd=eng.BASE,stdout=subprocess.PIPE,stderr=subprocess.STDOUT,text=True,bufsize=1,env=eng.engine_env(alias),start_new_session=True)
    expired=threading.Event(); killer=None
    def kill_group(sig):
        with contextlib.suppress(ProcessLookupError): os.killpg(proc.pid,sig)
    def stop():
        nonlocal killer
        expired.set();kill_group(signal.SIGTERM)
        killer=threading.Timer(5,lambda:kill_group(signal.SIGKILL));killer.daemon=True;killer.start()
    timer=threading.Timer(timeout,stop);timer.daemon=True;timer.start()
    result=dict(ok=False,limit_hit=False,budget_hit=False,timed_out=False,cost_usd=None,model_usage=None)
    terminal_error=False
    try:
        with eng.engine_log(jid) as log:
            for line in proc.stdout:
                log.write(line)
                try: ev=json.loads(line)
                except ValueError: continue
                if not isinstance(ev,dict): continue
                if eng.limit_event(ev): result['limit_hit']=True
                if ev.get('type')=='error': terminal_error=True
                if ev.get('type')=='result':
                    subtype=ev.get('subtype','')
                    result['budget_hit']=subtype in ('error_max_budget_usd','error_max_turns')
                    terminal_error=terminal_error or bool(ev.get('is_error')) or bool(ev.get('error')) or ev.get('status')=='error' or subtype.startswith('error')
                    result.update(cost_usd=ev.get('total_cost_usd'),model_usage=ev.get('modelUsage'),session_id=ev.get('session_id'))
                if ev.get('type')=='turn.failed': terminal_error=True
                if ev.get('type')=='assistant':
                    for c in ev.get('message',{}).get('content',[]):
                        if c.get('type')=='tool_use': eng.jset(jid,log_add='TOOL\t'+c.get('name',''))
            rc=proc.wait();log.write(f'\n--- bounded engine {alias}: rc={rc} ---\n')
        result.update(ok=rc==0 and not terminal_error and not expired.is_set(),rc=rc,timed_out=expired.is_set())
        return result
    finally:
        timer.cancel()
        if killer: killer.cancel()
        if expired.is_set() or proc.poll() is None: kill_group(signal.SIGKILL)
        proc.wait();proc.stdout.close()
